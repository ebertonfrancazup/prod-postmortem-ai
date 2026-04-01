import os
import json
import logging
import base64
import io
from datetime import datetime
import docx
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fastapi import HTTPException
import google.generativeai as genai

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, Flowable
from reportlab.lib.units import inch

from app.schemas import AnalyzeRequest, IncidentReport

logger = logging.getLogger(__name__)

api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
if api_key:
    genai.configure(api_key=api_key)
else:
    logger.warning("Neither GEMINI_API_KEY nor GOOGLE_API_KEY found in environment.")

async def process_incident_data(request: AnalyzeRequest) -> IncidentReport:
    """Sends incident data to the Gemini model optionally with images."""
    if not (os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")):
        raise HTTPException(status_code=500, detail="API Key not configured.")

    optional_context = []
    if request.ticket_number: optional_context.append(f"Ticket Number: {request.ticket_number}")
    if request.time_range: optional_context.append(f"Time Range: {request.time_range}")
    if request.affected_services: optional_context.append(f"Affected Services: {request.affected_services}")
    if request.impact: optional_context.append(f"Impact Classification: {request.impact}")
    if request.key_stakeholders: optional_context.append(f"Key Stakeholders: {request.key_stakeholders}")
    if request.customers: optional_context.append(f"Affected Customers: {request.customers}")
    context_str = "\n".join(optional_context) if optional_context else "None provided."

    images_instruction = ""
    if request.images:
        images_instruction = "IMPORTANT: You received screenshots (which will be organically attached to the final PDF). Read them to better understand the incident."
        
    time_range_instruction = ""
    if request.time_range:
        time_range_instruction = f"CRITICAL: The user explicitly defined the incident Time Range as '{request.time_range}'. This defines the boundaries of your analysis. IMPORTANT: `downtime_minutes` should be the actual time the service was degraded (Critical/Warning). If the logs show the system was healthy at the beginning and then crashed, downtime is only the crashed portion. HOWEVER, if the logs LACK explicit timestamps for when the failure started or ended, you MUST fully fallback to this '{request.time_range}' as the absolute truth for the start and end of the outage, assuming the entire window was degraded! Se a janela terminar e os logs não mostrarem recuperação, status é 'Ongoing'."

    # Load externalized prompt skill
    skill_path = os.path.join(os.path.dirname(__file__), "skills", "sre_report_skill.md")
    try:
        with open(skill_path, "r", encoding="utf-8") as f:
            prompt_template = f.read()
    except Exception as e:
        logger.error(f"Failed to load prompt skill at {skill_path}: {e}")
        # Severe fallback if file is missing
        prompt_template = "Analyze logs and return JSON. [LOGS]\n{logs}"

    prompt = prompt_template.replace("{logs}", request.logs)\
                            .replace("{transcription}", request.transcription or "None provided.")\
                            .replace("{context_str}", context_str)\
                            .replace("{images_instruction}", images_instruction)\
                            .replace("{time_range_instruction}", time_range_instruction)

    if request.language == "pt-br":
        prompt = "LANGUAGE INSTRUCTION: Generate ALL text content (titles, summaries, descriptions, timeline events, next steps) in Brazilian Portuguese (pt-BR). Do NOT use English for any content fields.\n\n" + prompt
    else:
        prompt = "LANGUAGE INSTRUCTION: Generate ALL text content in English.\n\n" + prompt

    try:
        model = genai.GenerativeModel('gemini-2.5-flash', generation_config={"response_mime_type": "application/json"})
        
        # Build multipart input
        contents = [prompt]
        for idx, b64_img in enumerate(request.images):
            if "," in b64_img:
                _, b64_img = b64_img.split(",", 1)
            try:
                decoded = base64.b64decode(b64_img)
                contents.append({
                    "mime_type": "image/jpeg", # Safe fallback
                    "data": decoded
                })
            except Exception as e:
                logger.error(f"Error decoding image: {e}")
                
        response = model.generate_content(contents)
        try:
            report_data = json.loads(response.text)
        except json.JSONDecodeError:
            import json_repair
            logger.warning("Standard JSON parse failed, utilizing json_repair.")
            report_data = json_repair.loads(response.text)
        
        # Capture AI Token Metrics
        if response.usage_metadata:
            report_data['token_usage'] = {
                "prompt_tokens": response.usage_metadata.prompt_token_count,
                "candidates_tokens": response.usage_metadata.candidates_token_count,
                "total_tokens": response.usage_metadata.total_token_count
            }
            
        return IncidentReport.model_validate(report_data)
    except Exception as e:
        logger.error(f"Error invoking Gemini: {e}")
        raise HTTPException(status_code=500, detail=str(e))

class ReportGenerator:
    def __init__(self, request: AnalyzeRequest):
        self.request = request
        self.styles = getSampleStyleSheet()

        # Extract StackSpot logo from Template.docx (DOCX is a ZIP archive)
        self._logo_data = None
        import zipfile
        tpl = os.path.join(os.path.dirname(__file__), 'templates', 'Template.docx')
        if os.path.exists(tpl):
            try:
                with zipfile.ZipFile(tpl, 'r') as z:
                    imgs = sorted([f for f in z.namelist() if f.startswith('word/media/')])
                    if imgs:
                        self._logo_data = z.read(imgs[0])
            except Exception as e:
                logger.warning(f"Could not extract logo from template: {e}")
        
        # Custom Styles
        self.styles.add(ParagraphStyle(name='HeaderLeft', parent=self.styles['Normal'], fontSize=16, leading=20, textColor=colors.whitesmoke, fontName='Helvetica-Bold'))
        self.styles.add(ParagraphStyle(name='HeaderRight', parent=self.styles['Normal'], fontSize=10, leading=14, textColor=colors.whitesmoke, alignment=2))
        self.styles.add(ParagraphStyle(name='MainTitle', parent=self.styles['Title'], fontSize=28, leading=34, textColor=colors.HexColor("#334155"), alignment=0, spaceAfter=2))
        self.styles.add(ParagraphStyle(name='SubTitle', parent=self.styles['Normal'], fontSize=16, leading=22, textColor=colors.HexColor("#4B5563"), spaceAfter=20))
        self.styles.add(ParagraphStyle(name='CardTitle', parent=self.styles['Normal'], fontSize=10, leading=12, textColor=colors.HexColor("#6B7280"), fontName='Helvetica-Bold'))
        self.styles.add(ParagraphStyle(name='CardValue', parent=self.styles['Normal'], fontSize=16, leading=20, textColor=colors.HexColor("#334155"), fontName='Helvetica-Bold', spaceBefore=6, spaceAfter=6))
        self.styles.add(ParagraphStyle(name='CardValueGreen', parent=self.styles['Normal'], fontSize=16, leading=20, textColor=colors.HexColor("#10B981"), fontName='Helvetica-Bold', spaceBefore=6, spaceAfter=6))
        self.styles.add(ParagraphStyle(name='CardValueSmall', parent=self.styles['Normal'], fontSize=13, leading=16, textColor=colors.HexColor("#334155"), fontName='Helvetica-Bold', spaceBefore=6, spaceAfter=6))
        self.styles.add(ParagraphStyle(name='CardValueSmallGreen', parent=self.styles['Normal'], fontSize=13, leading=16, textColor=colors.HexColor("#10B981"), fontName='Helvetica-Bold', spaceBefore=6, spaceAfter=6))
        self.styles.add(ParagraphStyle(name='CardValueMicro', parent=self.styles['Normal'], fontSize=10, leading=12, textColor=colors.HexColor("#334155"), fontName='Helvetica-Bold', spaceBefore=6, spaceAfter=6))
        self.styles.add(ParagraphStyle(name='CardValueMicroGreen', parent=self.styles['Normal'], fontSize=10, leading=12, textColor=colors.HexColor("#10B981"), fontName='Helvetica-Bold', spaceBefore=6, spaceAfter=6))
        self.styles.add(ParagraphStyle(name='CardSub', parent=self.styles['Normal'], fontSize=8, leading=10, textColor=colors.HexColor("#9CA3AF")))
        self.styles.add(ParagraphStyle(name='SectionTitle', parent=self.styles['Heading2'], fontSize=16, leading=20, textColor=colors.HexColor("#334155"), spaceAfter=10))
        self.styles.add(ParagraphStyle(name='ExecText', parent=self.styles['Normal'], fontSize=11, leading=16, textColor=colors.HexColor("#374151"), spaceAfter=14))
        self.styles.add(ParagraphStyle(name='SlaPercent', parent=self.styles['Normal'], fontSize=34, leading=40, textColor=colors.HexColor("#111827"), fontName='Helvetica-Bold', alignment=1))
        self.styles.add(ParagraphStyle(name='SlaLabel', parent=self.styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor("#6B7280"), fontName='Helvetica-Bold', alignment=1, spaceAfter=15))
        self.styles.add(ParagraphStyle(name='NextStepBullet', parent=self.styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor("#4B5563"), spaceAfter=8))
        self.styles.add(ParagraphStyle(name='TokenFooter', parent=self.styles['Normal'], fontSize=8, leading=10, textColor=colors.HexColor("#9CA3AF"), alignment=1))
        self.styles.add(ParagraphStyle(name='TimelineHeader', parent=self.styles['Normal'], fontSize=11, leading=14, textColor=colors.HexColor("#374151"), fontName='Helvetica-Bold'))
        self.styles.add(ParagraphStyle(name='TimelineText', parent=self.styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor("#374151")))

    def _header_footer(self, canvas, doc):
        canvas.saveState()

        # White header background
        canvas.setFillColor(colors.white)
        canvas.rect(0, 718, letter[0], 100, fill=1, stroke=0)

        # Orange accent line at the bottom of the header
        canvas.setFillColor(colors.HexColor("#ff6900"))
        canvas.rect(0, 718, letter[0], 3, fill=1, stroke=0)

        # StackSpot logo (extracted from Template.docx)
        if self._logo_data:
            from reportlab.lib.utils import ImageReader
            logo_reader = ImageReader(io.BytesIO(self._logo_data))
            lw, lh = logo_reader.getSize()
            target_h = 38
            target_w = target_h * (lw / lh)
            canvas.drawImage(logo_reader, 40, 740, width=target_w, height=target_h, mask='auto')
        else:
            canvas.setFillColor(colors.HexColor("#ff6900"))
            canvas.setFont("Helvetica-Bold", 20)
            canvas.drawString(40, 748, "StackSpot")

        # Right side: label + date
        canvas.setFillColor(colors.HexColor("#374151"))
        canvas.setFont("Helvetica-Bold", 11)
        canvas.drawRightString(letter[0] - 40, 762, "CONFIDENTIAL")
        canvas.setFont("Helvetica", 10)
        canvas.drawRightString(letter[0] - 40, 748, f"EXECUTIVE REPORT  |  {datetime.now().strftime('%B %d, %Y')}")

        canvas.restoreState()

    def generate_pdf(self, report: IncidentReport) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=85, bottomMargin=40)
        elements = []

        # Language-aware labels
        is_pt = self.request.language == "pt-br"
        L = {
            'exec_summary': 'Resumo Executivo' if is_pt else 'Executive Summary',
            'impact': 'Impacto' if is_pt else 'Impact',
            'root_cause': 'Causa Raiz' if is_pt else 'Root Cause',
            'resolution': 'Resolucao' if is_pt else 'Resolution',
            'next_steps': 'Proximos Passos' if is_pt else 'Suggested Next Steps',
            'timeline_title': 'Linha do Tempo da Resposta' if is_pt else 'Team Response Timeline',
            'timestamp': 'Horario' if is_pt else 'Timestamp',
            'event': 'Decisao / Evento' if is_pt else 'Decision / Event',
            'sla_budget': 'BUDGET SLA UTILIZADO' if is_pt else 'SLA BUDGET USED',
            'monitoring': 'Evidencia de Monitoramento' if is_pt else 'Monitoring Evidence',
            'threshold': 'Meta' if is_pt else 'Threshold Target',
        }

        elements.append(Paragraph("Post-Mortem", self.styles['MainTitle']))
        elements.append(Paragraph(report.metrics.incident_title, self.styles['SubTitle']))
        elements.append(Spacer(1, 10))

        # Cards - Modern Floating Style
        metrics = report.metrics
        usable_width = letter[0] - 80
        num_cards = 4
        gap = 10
        card_w = (usable_width - (gap * (num_cards - 1))) / num_cards

        def get_value_style(val_str, is_green):
            length = len(str(val_str))
            if length > 20:
                return self.styles['CardValueMicroGreen'] if is_green else self.styles['CardValueMicro']
            elif length > 12:
                return self.styles['CardValueSmallGreen'] if is_green else self.styles['CardValueSmall']
            return self.styles['CardValueGreen'] if is_green else self.styles['CardValue']

        class RoundedCard(Flowable):
            def __init__(self, title_p, sub_p, value_p, width, height, bg_color="#F1F5F9", outline_color="#E2E8F0"):
                Flowable.__init__(self)
                self.width = width
                self.height = height
                self.title_p = title_p
                self.sub_p = sub_p
                self.value_p = value_p
                self.bg_color = colors.HexColor(bg_color)
                self.outline_color = colors.HexColor(outline_color)

            def wrap(self, availWidth, availHeight):
                return self.width, self.height

            def draw(self):
                self.canv.saveState()
                self.canv.setFillColor(self.bg_color)
                self.canv.setStrokeColor(self.outline_color)
                self.canv.setLineWidth(0.5)
                self.canv.roundRect(0, 0, self.width, self.height, 16, fill=1, stroke=1)

                padding = 16
                available_w = self.width - padding * 2

                w_t, h_t = self.title_p.wrap(available_w, self.height)
                w_s, h_s = self.sub_p.wrap(available_w, self.height)
                w_v, h_v = self.value_p.wrap(available_w, self.height)

                y_cursor = self.height - padding

                self.title_p.drawOn(self.canv, padding, y_cursor - h_t)
                y_cursor -= (h_t + 2)

                # Only draw subtitle if it has actual content (h_s > 2 means non-empty)
                if h_s > 2:
                    self.sub_p.drawOn(self.canv, padding, y_cursor - h_s)
                    y_cursor -= (h_s + 8)
                else:
                    y_cursor -= 6

                # Clamp value height so it never overflows the card bottom
                max_v_h = max(y_cursor - padding, 10)
                if h_v > max_v_h:
                    w_v, h_v = self.value_p.wrap(available_w, max_v_h)
                self.value_p.drawOn(self.canv, padding, y_cursor - h_v)

                self.canv.restoreState()

        def create_modern_card(title, value, sub, is_green=False):
            val_style = get_value_style(value, is_green)
            title_p = Paragraph(title, self.styles['CardTitle'])
            sub_p = Paragraph(sub, self.styles['CardSub'])
            value_p = Paragraph(value, val_style)
            return RoundedCard(title_p, sub_p, value_p, width=card_w, height=105)

        cust_value = self.request.customers if self.request.customers else metrics.affected_customers
        cust_label = "Clientes Afetados" if (cust_value and (',' in cust_value or ' e ' in cust_value.lower())) else "Cliente Afetado"

        row1_cards = [
            create_modern_card("Criticidade", f"{metrics.impact}", ""),
            create_modern_card("Downtime / MTTR", f"{metrics.total_downtime}", ""),
            create_modern_card("Status", f"{metrics.service_status}", "", is_green=(metrics.service_status.lower()=="resolved")),
            create_modern_card(cust_label, f"{cust_value}", ""),
        ]

        def pack_cards(cards_list):
            r_cells = []
            col_widths = []
            for i, card in enumerate(cards_list):
                r_cells.append(card)
                col_widths.append(card_w)
                if i < len(cards_list) - 1:
                    r_cells.append("")
                    col_widths.append(gap)
            t = Table([r_cells], colWidths=col_widths)
            t.setStyle(TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('PADDING', (0, 0), (-1, -1), 0),
            ]))
            return t

        elements.append(pack_cards(row1_cards))
        elements.append(Spacer(1, 10))

        # Optional Row 2 — user-provided fields only
        row2_cards = []
        if self.request.slo:
            row2_cards.append(create_modern_card("SLO", self.request.slo, ""))
        if self.request.affected_requests_pct:
            row2_cards.append(create_modern_card("% Requisições Afetadas", self.request.affected_requests_pct, ""))
        if self.request.impacted_journeys:
            row2_cards.append(create_modern_card("Jornadas Impactadas", self.request.impacted_journeys, ""))
        if self.request.estimated_loss:
            row2_cards.append(create_modern_card("Perda Estimada", self.request.estimated_loss, ""))
            
        if row2_cards:
            for c in row2_cards:
                c._colWidths = [card_w]
            elements.append(pack_cards(row2_cards))
            elements.append(Spacer(1, 20))
        else:
            elements.append(Spacer(1, 10))

        # Two Columns Split (Executive Summary + SLA)
        sla_hours = self.request.sla_hours
        sla_max_minutes = sla_hours * 60
        actual_minutes = metrics.downtime_minutes
        
        sla_percent = min(100, round((actual_minutes / sla_max_minutes) * 100)) if sla_max_minutes > 0 else 0
        sla_color = "#10B981" if sla_percent < 50 else ("#F59E0B" if sla_percent < 80 else "#EF4444")
        
        left_col = [
            Paragraph(L['exec_summary'], self.styles['SectionTitle']),
            Paragraph(f"<b>{L['impact']}:</b> {report.executive_summary.impact}", self.styles['ExecText'])
        ]

        sla_circle_val = Paragraph(f"<font color='{sla_color}'>{sla_percent}%</font>", self.styles['SlaPercent'])
        right_col = [
            Spacer(1, 5),
            sla_circle_val,
            Paragraph(L['sla_budget'], self.styles['SlaLabel']),
            Paragraph(f"({sla_hours}h {L['threshold']})", self.styles['TokenFooter'])
        ]

        summary_table = Table([[left_col, right_col]], colWidths=[360, 160])
        summary_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ALIGN', (1, 0), (1, 0), 'CENTER'),
            ('BOX', (0, 0), (0, 0), 0.5, colors.HexColor("#E5E7EB")),
            ('BOX', (1, 0), (1, 0), 0.5, colors.HexColor("#E5E7EB")),
            ('PADDING', (0, 0), (-1, -1), 15)
        ]))

        elements.append(summary_table)
        elements.append(Spacer(1, 15))

        elements.append(Paragraph(f"<b>{L['root_cause']}:</b> {report.executive_summary.root_cause}", self.styles['ExecText']))
        elements.append(Paragraph(f"<b>{L['resolution']}:</b> {report.executive_summary.resolution}", self.styles['ExecText']))
        elements.append(Spacer(1, 15))

        elements.append(Paragraph(L['next_steps'], self.styles['SectionTitle']))
        for step in report.next_steps:
            elements.append(Paragraph(f"• {step}", self.styles['NextStepBullet']))

        elements.append(Spacer(1, 20))

        # Monitoring Evidence — user provided screenshots (after next steps, before timeline)
        if self.request.images:
            from reportlab.lib.utils import ImageReader
            for b64 in self.request.images:
                try:
                    if "," in b64:
                        _, b64 = b64.split(",", 1)
                    img_data = base64.b64decode(b64)
                    img_buffer = io.BytesIO(img_data)
                    rp_img = ImageReader(img_buffer)
                    iw, ih = rp_img.getSize()
                    aspect = ih / float(iw)
                    draw_w = min(usable_width, iw)
                    if iw > usable_width: draw_w = usable_width
                    draw_h = draw_w * aspect

                    elements.append(Paragraph(L['monitoring'], self.styles['SectionTitle']))
                    elements.append(RLImage(img_buffer, width=draw_w, height=draw_h))
                    elements.append(Spacer(1, 20))
                    break
                except Exception as e:
                    logger.error(f"Failed to embed user image: {e}")

        elements.append(Paragraph(L['timeline_title'], self.styles['SectionTitle']))
        timeline_data = [[
            Paragraph(L['timestamp'], self.styles['TimelineHeader']),
            Paragraph(L['event'], self.styles['TimelineHeader'])
        ]]
        for event in report.timeline:
            event_text = event.event
            if event.detail:
                event_text += f"<br/><font size='8' color='#6B7280'>{event.detail}</font>"
            timeline_data.append([
                Paragraph(event.timestamp, self.styles['TimelineText']),
                Paragraph(event_text, self.styles['TimelineText'])
            ])

        timeline_table = Table(timeline_data, colWidths=[120, 400])
        timeline_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(timeline_table)

        doc.build(elements, onFirstPage=self._header_footer, onLaterPages=self._header_footer)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

class DocxGenerator:
    def __init__(self, request: AnalyzeRequest):
        self.request = request

    def generate_docx(self, report: IncidentReport) -> bytes:
        import copy as copy_mod
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        template_path = os.path.join(os.path.dirname(__file__), 'templates', 'Template.docx')
        if os.path.exists(template_path):
            doc = Document(template_path)
        else:
            logger.warning("DOCX template not found, using default.")
            doc = Document()

        metrics = report.metrics
        customers = self.request.customers if self.request.customers else metrics.affected_customers

        def set_text(para, text):
            """Replace paragraph content keeping first run's formatting."""
            for run in para.runs:
                run.text = ''
            if para.runs:
                para.runs[0].text = str(text)
            else:
                para.add_run(str(text))

        def find_para(substr):
            for p in doc.paragraphs:
                if substr in p.text:
                    return p
            return None

        def find_next_empty(ref_para, lookahead=3):
            """Return the next empty paragraph after ref_para within lookahead distance."""
            found = False
            count = 0
            for p in doc.paragraphs:
                if p._p is ref_para._p:
                    found = True
                    continue
                if found:
                    if not p.text.strip():
                        return p
                    count += 1
                    if count >= lookahead:
                        break
            return None

        def insert_after(ref_para, text):
            """Insert a new paragraph after ref_para copying its style, return the new para."""
            new_p = copy_mod.deepcopy(ref_para._p)
            for r in new_p.findall(qn('w:r')):
                new_p.remove(r)
            ref_para._p.addnext(new_p)
            new_r = OxmlElement('w:r')
            new_t = OxmlElement('w:t')
            new_t.text = str(text)
            new_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            new_r.append(new_t)
            new_p.append(new_r)
            from docx.text.paragraph import Paragraph as DocxPara
            return DocxPara(new_p, new_p.getparent())

        # === HEADER FIELDS ===
        set_text(doc.paragraphs[1], metrics.incident_title)

        start_time = report.timeline[0].timestamp if report.timeline else ''
        p = find_para('Data/Hora do início do incidente:')
        if p:
            set_text(p, f"Data/Hora do início do incidente: {start_time}")

        end_time = report.timeline[-1].timestamp if report.timeline else ''
        p = find_para('Data/Hora do fim do incidente:')
        if p:
            set_text(p, f"Data/Hora do fim do incidente: {end_time}")

        p = find_para('Duração: 1 dia') or find_para('Duração:')
        if p:
            set_text(p, f"Duração: {metrics.total_downtime}")

        p = find_para('Status:')
        if p:
            set_text(p, f"Status: {metrics.service_status}")

        p = find_para('Impacto: ')
        if p:
            set_text(p, f"Impacto: {metrics.impact}")

        p = find_para('Impacto SLA:')
        if p:
            sla_max_min = self.request.sla_hours * 60
            sla_used = round((metrics.downtime_minutes / sla_max_min) * 100) if sla_max_min > 0 else 0
            sla_text = f"Impacto SLA: {sla_used}% do budget de SLA utilizado ({metrics.downtime_minutes}min de {sla_max_min}min)"
            if metrics.infra_slo:
                sla_text += f"\n\nImpacto SLO: {metrics.infra_slo}"
            set_text(p, sla_text)

        p = find_para('% de requisições afetadas:')
        if p:
            set_text(p, f"% de requisições afetadas: {self.request.affected_requests_pct or metrics.affected_users or ''}")

        p = find_para('Canal do incidente:')
        if p:
            set_text(p, f"Canal do incidente: {customers}")

        p = find_para('Perda estimada')
        if p:
            loss = self.request.estimated_loss or ''
            proactive = 'Sim' if self.request.proactive_incident else 'Nao'
            set_text(p, f"Perda estimada (receita ou produtividade): {loss}\n\nIncidente Proativo? {proactive}")

        # Clear the secondary "Duração: 1 hora" placeholder line
        p = find_para('Duração: 1 hora')
        if p:
            set_text(p, '')

        # === EXECUTIVE SUMMARY SECTION ===
        p = find_para('Resumo executivo do Incidente:')
        if p:
            set_text(p, f"📋 Resumo executivo do Incidente: {metrics.incident_title}")
            empty = find_next_empty(p)
            if empty:
                set_text(empty, report.executive_summary.impact)

        p = find_para('Causa Raiz (Root Cause)')
        if p:
            empty = find_next_empty(p)
            if empty:
                set_text(empty, report.executive_summary.root_cause)

        p = find_para('Impactos Principais')
        if p:
            empty = find_next_empty(p)
            if empty:
                set_text(empty, report.executive_summary.impact)

        p = find_para('Identificação:')
        if p:
            set_text(p, f"Identificação: {report.executive_summary.root_cause}")

        p = find_para('Ação Imediata:')
        if p:
            set_text(p, f"Ação Imediata: {report.executive_summary.impact}")

        p = find_para('Resolução Definitiva:')
        if p:
            set_text(p, f"Resolução Definitiva: {report.executive_summary.resolution}")

        # === NEXT STEPS ===
        p = find_para('Plano de Governança')
        if p and report.next_steps:
            set_text(p, report.next_steps[0])
            current = p
            for step in report.next_steps[1:]:
                current = insert_after(current, step)

        # === MONITORING EVIDENCE (images) ===
        if self.request.images:
            for b64 in self.request.images:
                try:
                    if "," in b64:
                        _, b64 = b64.split(",", 1)
                    img_data = base64.b64decode(b64)
                    img_stream = io.BytesIO(img_data)
                    doc.add_heading('Evidencia de Monitoramento', level=2)
                    doc.add_picture(img_stream, width=Inches(6.0))
                    doc.add_paragraph()
                    break
                except Exception as e:
                    logger.error(f"Failed to embed image in DOCX: {e}")

        # === VISÃO DETALHADA ===
        if report.timeline:
            detection_kws = ['detect', 'identif', 'aciona', 'alerta', 'recebid', 'report']
            mitigation_kws = ['mitiga', 'parcial', 'paliativ', 'workaround', 'contorn']
            resolution_kws = ['resolv', 'deploy', 'fix', 'corrig', 'encerr', 'restabelec']

            p = find_para('IDENTIFICAÇÃO/DETECÇÃO DETALHADA')
            if p:
                ev = next((e for e in report.timeline if any(k in e.event.lower() for k in detection_kws)), report.timeline[0])
                set_text(p, f"IDENTIFICAÇÃO/DETECÇÃO DETALHADA | {ev.timestamp}")
                empty = find_next_empty(p)
                if empty:
                    set_text(empty, ev.event)

            p = find_para('RESOLUÇÃO PARCIAL/MITIGAÇÃO')
            if p:
                ev = next((e for e in report.timeline if any(k in e.event.lower() for k in mitigation_kws)), None)
                if ev:
                    set_text(p, f"RESOLUÇÃO PARCIAL/MITIGAÇÃO | {ev.timestamp}")
                    empty = find_next_empty(p)
                    if empty:
                        set_text(empty, ev.event)

            p = find_para('RESOLUÇÃO |')
            if p:
                ev = next((e for e in reversed(report.timeline) if any(k in e.event.lower() for k in resolution_kws)), report.timeline[-1])
                set_text(p, f"RESOLUÇÃO | {ev.timestamp}")
                empty = find_next_empty(p)
                if empty:
                    set_text(empty, ev.event)

            p = find_para('RESOLUÇÃO DETALHADA')
            if p:
                empty = find_next_empty(p)
                if empty:
                    set_text(empty, report.executive_summary.resolution)

        # === TIMELINE ===
        p = find_para('Engajamento do Time IDP')
        if p and report.timeline:
            first = report.timeline[0]
            first_text = f"{first.timestamp} - {first.event}"
            if first.detail:
                first_text += f"\n{first.detail}"
            set_text(p, first_text)
            current = p
            for event in report.timeline[1:]:
                evt_text = f"{event.timestamp} - {event.event}"
                if event.detail:
                    evt_text += f"\n{event.detail}"
                current = insert_after(current, evt_text)
        elif p:
            set_text(p, '')

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
